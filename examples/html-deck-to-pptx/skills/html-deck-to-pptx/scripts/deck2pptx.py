#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""deck2pptx.py —— HTML deck 的版面 JSON → 原生「可编辑」PPTX。

为什么不是截图贴图：
    客户要的是**可编辑** PPT —— 卡片、色块、发丝线、表格线、图片都是 PowerPoint
    形状与图片，所有文字都是真正的文本框（run 级保留加粗/字号/颜色），位置与 HTML
    一一对应，打开后可以直接改字、换图、调版式。

输入：extract-deck.mjs 产出的 _layout.json（几何 + 文字 + 图片，全部取自浏览器 DOM）
输出：与 HTML 同名的 .pptx（默认放在 layout 文件旁边）

用法：
    python3 deck2pptx.py <_layout.json> [--out out.pptx] [--img-base DIR]
                        [--title 材料名] [--check] [--no-notes]
                        [--nowrap-classes a,b,c] [--scale-pt 1.0]
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import struct
import subprocess
import sys
import tempfile
import time
import urllib.parse

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

# PPT 画布永远是 13.333in 宽（16:9 的常规 PPT 尺寸），设计 px 按比例换算
CANVAS_EMU = 12192000
CANVAS_PT = 960.0            # 13.333in × 72pt/in

LATIN_FONT = "Arial"
EA_FONT = "PingFang SC"

# python-pptx 只认这些位图格式（PowerPoint 本身也认 SVG，但 python-pptx 写不进去）
SUPPORTED_EXT = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".wmf", ".emf"}
CACHE_DIR = os.path.join(tempfile.gettempdir(), "deck2pptx-cache")
CHROME = os.environ.get("CHROME_BIN") or "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"


# ────────────────────── 图片格式兼容（svg / webp → png） ──────────────────────

def png_size(path):
    """读 PNG 头拿像素尺寸（不做完整解码）。"""
    try:
        with open(path, "rb") as f:
            head = f.read(24)
        if head[:8] == b"\x89PNG\r\n\x1a\n" and head[12:16] == b"IHDR":
            return struct.unpack(">II", head[16:24])
    except OSError:
        pass
    return 0, 0


def _run_chrome(args, out, timeout=25.0):
    """Chrome 打完文件经常不退进程 → 后台起 + 轮询文件出现且稳定 + 主动收掉。"""
    proc = subprocess.Popen(args, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        t0, prev, stable = time.time(), -1, 0
        while time.time() - t0 < timeout:
            time.sleep(0.2)
            size = os.path.getsize(out) if os.path.exists(out) else 0
            stable = stable + 1 if (size and size == prev) else 0
            prev = size
            if stable >= 3:
                return True
        return os.path.exists(out)
    finally:
        proc.terminate()
        try:
            proc.wait(timeout=3)
        except subprocess.TimeoutExpired:
            proc.kill()


def _render_via_chrome(src_path, out_png, box_w, box_h):
    """用 Chrome 把图按「展示框的比例」渲成 PNG（SVG 只能这么栅格化）。

    为什么用 HTML 包一层：直接打开 .svg 会带文档默认边距、并按视口缩放；
    包一层 100%×100% + object-fit:fill 才与 deck 里的显示完全一致。
    """
    if not os.path.exists(CHROME):
        return False
    w = max(32, min(int(round(box_w * 2)), 2400))
    h = max(32, min(int(round(box_h * 2)), 2400))
    os.makedirs(CACHE_DIR, exist_ok=True)
    wrapper = os.path.join(CACHE_DIR, "wrap-" + hashlib.md5(out_png.encode()).hexdigest()[:10] + ".html")
    href = "file://" + urllib.parse.quote(os.path.abspath(src_path))
    with open(wrapper, "w", encoding="utf-8") as f:
        f.write("<!doctype html><meta charset=utf-8>"
                "<style>html,body{margin:0;padding:0;background:transparent;overflow:hidden}"
                "img{display:block;width:100%;height:100%;object-fit:fill}</style>"
                f'<img src="{href}">')
    return _run_chrome([CHROME, "--headless", "--disable-gpu", "--hide-scrollbars",
                        "--no-first-run", "--no-default-browser-check",
                        "--default-background-color=00000000",
                        f"--screenshot={out_png}", f"--window-size={w},{h}",
                        f"--user-data-dir={os.path.join(CACHE_DIR, 'profile')}",
                        "file://" + urllib.parse.quote(wrapper)], out_png)


def convert_image(path, box_w, box_h):
    """返回 (可用路径, 宽, 高)；宽高为 None 表示沿用原始 intrinsic 尺寸。"""
    ext = os.path.splitext(path)[1].lower()
    if ext in SUPPORTED_EXT:
        return path, None, None
    os.makedirs(CACHE_DIR, exist_ok=True)
    key = f"{os.path.abspath(path)}|{round(box_w)}|{round(box_h)}|{os.path.getmtime(path)}"
    out = os.path.join(CACHE_DIR, f"{os.path.splitext(os.path.basename(path))[0]}-"
                                  f"{hashlib.md5(key.encode()).hexdigest()[:10]}.png")
    if not os.path.exists(out):
        ok = False
        if ext != ".svg" and shutil.which("sips"):        # webp/heic/avif：sips 保真
            ok = subprocess.run(["sips", "-s", "format", "png", path, "--out", out],
                                capture_output=True).returncode == 0 and os.path.exists(out)
        if not ok:                                        # svg（或 sips 失败）：走 Chrome
            ok = _render_via_chrome(path, out, box_w, box_h)
        if not ok:
            print(f"  ! 图片格式不支持又转换失败：{os.path.basename(path)}（{ext}）")
            return None, 0, 0
        print(f"  · 已转换 {os.path.basename(path)} → PNG（{ext} → .png）")
    w, h = png_size(out)
    return out, (w or int(box_w)), (h or int(box_h))

ALIGN = {"start": PP_ALIGN.LEFT, "left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
         "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}


# ────────────────────────── 载入 / 工具 ──────────────────────────

def load_layout(path: str) -> dict:
    """容忍三种形状：{design,slides} / 页数组 / 旧的 `Result: "<json>"` 包装。"""
    raw = open(path, encoding="utf-8").read()
    if "Result: " in raw and not raw.lstrip().startswith("{"):
        seg = raw.split("Result: ", 1)[1].strip()
        i, j = seg.index('"'), seg.rindex('"')
        raw = json.loads(seg[i:j + 1])
    data = json.loads(raw)
    if isinstance(data, list):
        return {"design": [1920, 1080], "slides": data, "source": ""}
    return data


def rgb(css):
    if not css or css == "none":
        return None
    css = css.strip()
    if css.startswith("rgba") or css.startswith("hsla"):
        nums = css[css.index("(") + 1:css.index(")")].split(",")
        if len(nums) == 4 and float(nums[3]) == 0:
            return None
        return RGBColor(*[int(float(v)) for v in nums[:3]])
    if css.startswith("rgb"):
        nums = css[css.index("(") + 1:css.index(")")].split(",")[:3]
        return RGBColor(*[int(float(v)) for v in nums])
    return None


def px_of(v) -> float:
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace("px", "").strip() or 0)
    except ValueError:
        return 0.0


class Metrics:
    def __init__(self, design, scale_pt=1.0):
        self.dw, self.dh = design
        self.emu = CANVAS_EMU / self.dw
        self.pt = (CANVAS_PT / self.dw) * scale_pt

    def E(self, px) -> Emu:
        return Emu(int(round(px * self.emu)))

    def PT(self, px) -> Pt:
        return Pt(round(px * self.pt, 1))


# ────────────────────────── 画形状 ──────────────────────────

def add_rect(slide, M, r, fill, line, radius_px=0, line_w=0.75):
    if r[2] <= 0 or r[3] <= 0:
        return None
    rounded = radius_px >= (r[3] / 2 - 1) and radius_px > 6
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        M.E(r[0]), M.E(r[1]), M.E(r[2]), M.E(r[3]))
    if rounded:
        shape.adjustments[0] = 0.5
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_w)
    shape.shadow.inherit = False
    return shape


def add_bar(slide, M, r, color):
    """细线 / 竖条：用无边框实心矩形画，避免 PPT 直线的连接点差异。"""
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, M.E(r[0]), M.E(r[1]),
                               M.E(max(r[2], 0.6)), M.E(max(r[3], 0.6)))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def style_run(run, size_px, bold, color, M):
    f = run.font
    f.size = M.PT(size_px)
    f.bold = bold
    f.color.rgb = color
    f.name = LATIN_FONT
    # a:rPr 的子元素顺序必须符合 schema（a:latin 已由 font.name 建好）→ 只能 append ea/cs
    rPr = run._r.get_or_add_rPr()
    for tag, face in (("a:ea", EA_FONT), ("a:cs", LATIN_FONT)):
        if rPr.find(qn(tag)) is None:
            etree.SubElement(rPr, qn(tag)).set("typeface", face)


def add_text(slide, M, item, nowrap_kinds):
    r = item["r"]
    pad = item.get("pad") or [0, 0, 0, 0]          # [上, 右, 下, 左]
    x, y = r[0] + pad[3], r[1] + pad[0]
    w = max(r[2] - pad[1] - pad[3], 8)
    h = max(r[3] - pad[0] - pad[2], 8)

    lh = item.get("lh") or 0
    kinds = set(item.get("cls") or []) | {item.get("k", ""), item.get("tag", "")}
    nowrap = "nowrap" in (item.get("ws") or "")
    single = nowrap or bool(kinds & nowrap_kinds) or (lh > 0 and h <= lh * 1.6)
    # 单行元素：不给自动换行（字体度量差异会让它多出一行）；多行元素给一点余量
    box_h = min(h + (2 if single else 6), M.dh - y)

    box = slide.shapes.add_textbox(M.E(x), M.E(y), M.E(w), M.E(box_h))
    tf = box.text_frame
    tf.word_wrap = not single
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if item.get("va") == "middle" else MSO_ANCHOR.TOP

    align = ALIGN.get(item.get("ta"), PP_ALIGN.LEFT)
    base_size = item.get("fs") or 18
    base_bold = (item.get("fw") or 400) >= 600
    base_color = rgb(item.get("co")) or RGBColor(0x1B, 0x19, 0x15)

    # runs 里的 \n → 拆成新段落（PPT 里段落比软换行更好编辑）
    # 注意：只丢「空串」，不要用 strip() 判空 —— Python 的 strip 会把全角空格 U+3000
    # 也当空白，含独立全角空格的 run 会被误删（校对的「文字缺失」就是这么来的）。
    paragraphs: list[list] = [[]]
    for run in item.get("runs") or []:
        for i, part in enumerate(str(run.get("t", "")).split("\n")):
            if i:
                paragraphs.append([])
            if part != "":
                paragraphs[-1].append({**run, "t": part})

    for pi, runs in enumerate(paragraphs):
        p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        p.alignment = align
        if lh:
            p.line_spacing = M.PT(lh)
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        for run in runs:
            rr = p.add_run()
            rr.text = run["t"]
            style_run(rr,
                      run.get("fs") or base_size,
                      bool(run.get("b")) or (run.get("fw", 400) >= 600) or base_bold,
                      rgb(run.get("co")) or base_color, M)
    return box, single


def add_picture(slide, M, item, img_base):
    src = str(item.get("src") or "")
    path = src if os.path.isabs(src) else os.path.normpath(os.path.join(img_base, src))
    if not src or not os.path.exists(path):
        print(f"  ! 图片缺失：{path or src}")
        return None
    r = item["r"]
    if r[2] <= 0 or r[3] <= 0:
        return None
    # SVG / WebP 等 python-pptx 写不进去的格式 → 先转成 PNG（按展示框比例栅格化）
    path, cw, ch = convert_image(path, r[2], r[3])
    if path is None:
        return None
    nw = cw or item.get("nw") or 0
    nh = ch or item.get("nh") or 0
    fit = item.get("fit") or "fill"
    clip = item.get("crop") or [0, 0, 0, 0]           # [左, 右, 上, 下] 相对元素显示框的比例
    if not any(clip):
        if fit in ("contain", "none", "scale-down") and nw and nh:
            scale = min(r[2] / nw, r[3] / nh)
            w, h = nw * scale, nh * scale
            return slide.shapes.add_picture(path, M.E(r[0] + (r[2] - w) / 2),
                                            M.E(r[1] + (r[3] - h) / 2), M.E(w), M.E(h))
        pic = slide.shapes.add_picture(path, M.E(r[0]), M.E(r[1]), M.E(r[2]), M.E(r[3]))
    else:
        pic = slide.shapes.add_picture(path, M.E(r[0]), M.E(r[1]), M.E(r[2]), M.E(r[3]))

    # object-fit:cover 的基础裁剪 + 祖先 overflow:hidden 造成的裁切，按源图比例叠加
    cl = cr = ct = cb = 0.0
    if fit == "cover" and nw and nh:
        target, source = r[2] / r[3], nw / nh
        if source > target:
            cl = cr = (1 - target / source) / 2
        elif source < target:
            ct = cb = (1 - source / target) / 2
    if any(clip):
        sx, sy = 1 - cl - cr, 1 - ct - cb
        cl, cr = cl + clip[0] * sx, cr + clip[1] * sx
        ct, cb = ct + clip[2] * sy, cb + clip[3] * sy
    pic.crop_left, pic.crop_right, pic.crop_top, pic.crop_bottom = cl, cr, ct, cb
    return pic


# ────────────────────────── 主流程 ──────────────────────────

def build(layout: dict, img_base: str, title: str, notes: bool, nowrap_kinds: set,
          scale_pt: float) -> Presentation:
    M = Metrics(layout.get("design") or [1920, 1080], scale_pt)
    slides_data = layout["slides"]
    prs = Presentation()
    prs.slide_width = Emu(CANVAS_EMU)
    prs.slide_height = Emu(int(round(CANVAS_EMU * M.dh / M.dw)))
    blank = prs.slide_layouts[6]

    for i, page in enumerate(slides_data, 1):
        slide = prs.slides.add_slide(blank)
        add_rect(slide, M, [0, 0, M.dw, M.dh], rgb(page.get("bg")), None)

        for b in page.get("boxes") or []:
            fill = rgb(b.get("bg"))
            bgi = b.get("bgi") or "none"
            if fill is None and bgi != "none":          # 渐变 → 取渐变色带里的第一段橙
                fill = RGBColor(0xE8, 0x59, 0x0C)
            sides = [b.get("bt"), b.get("bl"), b.get("bb"), b.get("br")]   # 上/左/下/右
            widths = [px_of((s or [0])[0]) for s in sides]
            cols = [rgb((s or [None, None])[1]) for s in sides]
            line = None
            if widths[0] > 0 and len({(w, c) for w, c in zip(widths, cols)}) == 1:
                line = cols[0]                          # 四边同色同宽 → 统一边框
            add_rect(slide, M, b["r"], fill, line, px_of(b.get("rad")))
            if line is None:                            # 逐边描边 → 用细条补
                x, y, w, h = b["r"]
                if widths[1] > 0 and cols[1]:
                    add_bar(slide, M, [x, y, widths[1], h], cols[1])
                if widths[3] > 0 and cols[3]:
                    add_bar(slide, M, [x + w - widths[3], y, widths[3], h], cols[3])
                if widths[0] > 0 and cols[0]:
                    add_bar(slide, M, [x, y, w, widths[0]], cols[0])
                if widths[2] > 0 and cols[2]:
                    add_bar(slide, M, [x, y + h - widths[2], w, widths[2]], cols[2])

        for im in page.get("imgs") or []:
            add_picture(slide, M, im, img_base)

        for t in page.get("texts") or []:
            add_text(slide, M, t, nowrap_kinds)

        if notes:
            head = f"《{title}》第 {i} / {len(slides_data)} 页" if title else f"第 {i} / {len(slides_data)} 页"
            body = [n for n in (page.get("notes") or []) if n.strip()]
            slide.notes_slide.notes_text_frame.text = "\n".join([head] + body)
    return prs


def check(prs: Presentation, design) -> int:
    sw, sh = prs.slide_width, prs.slide_height
    bad = 0
    for i, slide in enumerate(prs.slides, 1):
        texts = 0
        for s in slide.shapes:
            if s.has_text_frame and s.text_frame.text.strip():
                texts += 1
            if s.left is None:
                continue
            if (s.left < -1000 or s.top < -1000
                    or s.left + (s.width or 0) > sw + 1000 or s.top + (s.height or 0) > sh + 1000):
                print(f"  ! 第 {i} 页越界：{s.shape_type} @({s.left},{s.top},{s.width},{s.height})")
                bad += 1
        print(f"  第 {i} 页：形状 {len(slide.shapes)} · 文本框 {texts}"
              f" · 备注 {len(slide.notes_slide.notes_text_frame.text)} 字")
    print("越界形状：", bad)
    return bad


def main() -> int:
    ap = argparse.ArgumentParser(description="HTML deck 版面 JSON → 原生可编辑 PPTX")
    ap.add_argument("layout")
    ap.add_argument("--out")
    ap.add_argument("--img-base")
    ap.add_argument("--title", default="")
    ap.add_argument("--no-notes", action="store_true")
    ap.add_argument("--nowrap-classes", default="")
    ap.add_argument("--scale-pt", type=float, default=1.0, help="字号整体缩放（默认 1.0）")
    ap.add_argument("--check", action="store_true", help="只自检，不写文件")
    a = ap.parse_args()

    layout = load_layout(a.layout)
    src = layout.get("source") or ""
    img_base = a.img_base or (os.path.dirname(src) if src else os.path.dirname(os.path.abspath(a.layout)))
    title = a.title or layout.get("title") or ""
    nowrap = {s.strip() for s in a.nowrap_classes.split(",") if s.strip()}

    print(f"载入 {len(layout['slides'])} 页布局 · 画布 {layout.get('design') or [1920, 1080]}"
          f" · 图片基准目录 {img_base}")
    prs = build(layout, img_base, title, not a.no_notes, nowrap, a.scale_pt)
    bad = check(prs, layout.get("design"))

    if not a.check:
        out = a.out
        if not out:
            base = os.path.splitext(src or a.layout)[0] or "deck"
            out = base + ".pptx"
        prs.save(out)
        print("已写出：", out, f"({os.path.getsize(out) / 1024:.0f} KB)")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
